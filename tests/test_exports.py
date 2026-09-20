"""Synthetic tests for the two exporters — no real statement data anywhere.

An export is the only artifact a buyer actually opens, so it must be tested the
same way a gate is: with data built here, so the test runs in CI (where the
629-page cache does not exist) and cannot pass by accident on one machine.

What is checked, and why each one matters:
  · sheet names/count — a missing sheet is the classic silent loss;
  · the printed value and the parsed number both survive — the whole point of
    the export is that it never shows one without the other;
  · an unproven page reaches «ما لم يُثبت» with its reason — a report that hides
    its gaps is the failure this project exists to prevent;
  · the markdown → HTML converter keeps RTL and real table rows.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

PROJ = Path(__file__).resolve().parents[1]
for _p in (str(PROJ), str(PROJ / "src")):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from openpyxl import load_workbook  # noqa: E402

from tools.to_pdf import render  # noqa: E402
PROFILE = Path(__file__).resolve().parents[1] / "profiles" / "al-rajhi.json"
from tools.to_xlsx import build  # noqa: E402


def _synthetic_run(tmp_path: Path) -> Path:
    """A two-page run: page 1 verified, page 2 has no printable totals.

    Dates are deliberately mixed the way the real paper is: Arabic-Indic digits,
    Extended Arabic-Indic (Persian) digits, and one value mixing both sets.
    """
    run = tmp_path / "run"
    (run / "results").mkdir(parents=True)
    (run / "results" / "pg-001.json").write_text(json.dumps({
        "pg": 1, "page_no": 1,
        "raw_rows": [
            {"movement": "300.00", "balance": "100.00", "desc": "تحويل",
             "date": "٢٠٢٤٠١٠١", "raw_movement": "٣٠٠,٠٠", "raw_balance": "١٠٠,٠٠"},
            {"movement": "50.00", "balance": "50.00", "desc": "سحب",
             "date": "۲۰۱۳۱۲۰٥", "raw_movement": "٥٠,٠٠", "raw_balance": "٥٠,٠٠"},
        ],
        "footer": {"debits": "300.00", "credits": "0.00", "balance": "100.00"},
        "arbitrated_by": [{"field": "balance", "why": "test"}],
    }, ensure_ascii=False), encoding="utf-8")
    (run / "results" / "pg-002.json").write_text(json.dumps({
        "pg": 2, "page_no": 2, "raw_rows": [], "footer": None,
    }, ensure_ascii=False), encoding="utf-8")
    (run / "slice_report.json").write_text(json.dumps({
        "slice": {"first": 1, "count": 2, "pages_done": 2},
        "totals": {"rows": 2, "clean": 2, "clean_ratio": 1.0},
        "footer_role": "cumulative_printed_totals",
        "usage": {"cost": 0.0123, "calls": 2},
        "footer": {"ok": 1, "mismatch": 0, "absent": 1, "unchecked": 0, "gap": 0},
        "page_numbers": {"gaps": []},
        "per_page": [
            {"page": 1, "rows": 2, "footer": "ok", "suspects": 0, "page_no": 1,
             "origin": "live", "ms_read": 1000,
             "footer_detail": [{"field": "debits", "footer": "300.00", "ok": True},
                               {"field": "credits", "footer": "0.00", "ok": True},
                               {"field": "balance", "footer": "100.00", "ok": True}]},
            {"page": 2, "rows": 0, "footer": "absent", "suspects": 0, "page_no": 2,
             "origin": "live", "ms_read": 2000},
        ],
    }, ensure_ascii=False), encoding="utf-8")
    return run


def test_xlsx_has_the_seven_sheets(tmp_path):
    out = tmp_path / "export.xlsx"
    info = build(_synthetic_run(tmp_path), out, None, profile=PROFILE)
    wb = load_workbook(out)
    assert wb.sheetnames == ["الملخص", "الحركات", "الأسئلة", "صفوف الأسئلة",
                             "التحقق لكل صفحة", "ما لم يُثبت", "كيف تُقرأ هذه الأوراق"]
    assert info["rows"] == 2 and info["pages"] == 2


def test_dates_are_normalised_across_all_three_digit_sets():
    """The defect a buyer sees first: one column, three numeral systems, mixed."""
    from tools.to_xlsx import normalize_date

    assert normalize_date("٢٠١٣١٠٣١")[:2] == ("2013-10-31", "ok")      # عربية-هندية
    assert normalize_date("۲۰۱۳۱۱۲۹")[:2] == ("2013-11-29", "ok")      # فارسية
    assert normalize_date("2016/08/08")[:2] == ("2016-08-08", "ok")    # لاتينية بفواصل
    assert normalize_date("۲۰۱۳۱۲۰٥")[:2] == ("2013-12-05", "ok")      # مختلطة داخلها
    assert normalize_date(None)[:2] == ("", "missing")
    assert normalize_date("٥")[:2] == ("", "incomplete")
    assert normalize_date("٢٠٢٤١٣٠١")[:2] == ("", "implausible")        # شهر ١٣


def test_xlsx_keeps_both_the_printed_and_the_parsed_number(tmp_path):
    """A column a human must trust (exactly what the paper printed) beside one a
    filter can use (parsed). The parsed column now carries the value the balance
    chain proved, not the raw read — both stay visible."""
    out = tmp_path / "export.xlsx"
    build(_synthetic_run(tmp_path), out, None, profile=PROFILE)
    ws = load_workbook(out)["الحركات"]
    header = [c.value for c in ws[1]]
    row = [c.value for c in ws[2]]
    assert "الحركة كما طُبعت" in header and "الحركة المثبتة بالسلسلة" in header
    assert "مدين" in header and "دائن" in header        # عمودا البنك الموجبان
    assert "حكم السلسلة على الصفّ" in header
    assert row[header.index("الحركة كما طُبعت")] == "٣٠٠,٠٠"     # ما على الورقة (نصّ)
    # المثبت رقم لا نصّ: يُفرَّز ويُجمع. القيمة 300.00 تُخزَّن 300.
    assert float(str(row[header.index("الحركة المثبتة بالسلسلة")])) == 300.0
    assert row[header.index("التاريخ (ميلادي)")] == "2024-01-01"
    assert row[header.index("حالة التاريخ")] == "تاريخ كامل"


def test_summary_counts_the_date_statuses_including_failures(tmp_path):
    """A summary that hides the bad dates behind an empty cell is not a summary."""
    out = tmp_path / "export.xlsx"
    build(_synthetic_run(tmp_path), out, None, profile=PROFILE)
    ws = load_workbook(out)["الملخص"]
    facts = {str(r[0]): r[1] for r in ws.iter_rows(min_row=2, values_only=True)}
    assert facts["تاريخ كامل (٨ خانات)"] == 2
    assert facts["الصفحات"] == 2
    assert facts["الصفحات المطابقة لإجمالياتها المطبوعة"] == 1


def test_questions_sheet_is_deterministic_and_names_its_method(tmp_path):
    out = tmp_path / "export.xlsx"
    build(_synthetic_run(tmp_path), out, None, profile=PROFILE)
    wb = load_workbook(out)
    ws = wb["الأسئلة"]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    assert len(rows) >= 10
    assert all(r[3] and "نموذج" not in str(r[3]) for r in rows), "الطريقة يجب أن تُذكر"
    biggest = next(r for r in rows if "أكبر حركة" in str(r[1]))
    assert biggest[2] == "300.00"
    ev = wb["صفوف الأسئلة"]
    ev_rows = [[c.value for c in r] for r in ev.iter_rows(min_row=2)]
    assert any(r[0] == 7 and r[1] == 1 for r in ev_rows), "الصفّ الداعم لأكبر حركة"


def test_xlsx_names_the_unproven_page_with_its_reason(tmp_path):
    out = tmp_path / "export.xlsx"
    build(_synthetic_run(tmp_path), out, None, profile=PROFILE)
    ws = load_workbook(out)["ما لم يُثبت"]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    assert len(rows) == 1
    assert rows[0][0] == 2
    assert "لا سطر إجماليات مطبوع" in str(rows[0][3])


def test_a_scan_gap_names_the_missing_printed_sheets_not_a_file_page(tmp_path):
    """فجوة المسح لا تتّهم صفحةً في الملف — لأن الموقع ليس هوية.

    بنية المُنتِج `(موقع_قبل, مطبوع_قبل, موقع_بعد, مطبوع_بعد, [أرقام غائبة])`:
    الموقع 426 في الملف هو المطبوعة 425 وهي **مقروءة ومحسوبة**، فوضعُ الموقع في
    عمود «الصفحة (ملف)» يتّهم صفحةً موجودة، ووضعُه في عمود «رقم الصفحة» يُخفي
    أيّ الأوراق غاب. القاعدة: العمود الأول فارغ، والثاني يحمل الغائب المطبوع.
    """
    run = _synthetic_run(tmp_path)
    report = json.loads((run / "slice_report.json").read_text(encoding="utf-8"))
    report["page_numbers"] = {"gaps": [[1, 0, 2, 3, [1, 2]]]}
    (run / "slice_report.json").write_text(json.dumps(report, ensure_ascii=False),
                                          encoding="utf-8")
    out = tmp_path / "export.xlsx"
    build(run, out, None, profile=PROFILE)
    rows = [[c.value for c in r]
            for r in load_workbook(out)["ما لم يُثبت"].iter_rows(min_row=2)]
    gap = next(r for r in rows if "فجوة مسح" in str(r[3]))
    assert gap[0] is None, "لا صفحة ملف تُتَّهم بالفجوة"
    assert gap[1] == "1 · 2", "أرقام الأوراق الغائبة وحدها"
    assert "بين الملفين 1 و2" in str(gap[3])
    # وصياغة **قيد الفجوة** في ورقة «الحركات» تصف «ما بين الرصيدين» لا «حركات
    # الأوراق الغائبة» وحدها — وتُقاس على الملف المبني (proof: الرسالة الرابعة)،
    # لأن هذا الـfixture يبني فجوةً بلا تذييل مطبوع قبلها فلا يُثبت قيدها.


def test_movement_source_column_separates_the_chain_from_the_paper(tmp_path):
    """عمود «المثبتة بالسلسلة» لا يكفي: في مرساة الفجوة يحمل مبلغاً مطبوعاً.

    رفعه مدقّق خارجي: قيمةٌ من الورق في عمودٍ اسمه السلسلة. فالعلاج صنفٌ — عمود
    «مصدر إثبات الحركة» يُعلن مصدر كل حركة، ويُنبَّه إن خالف إعلانُه حكمَ السلسلة.
    """
    out = tmp_path / "export.xlsx"
    build(_synthetic_run(tmp_path), out, None, profile=PROFILE)
    ws = load_workbook(out)["الحركات"]
    header = [c.value for c in ws[1]]
    assert "مصدر إثبات الحركة" in header
    assert header.index("مصدر إثبات الحركة") == header.index("الحركة المثبتة بالسلسلة") + 1
    i = header.index("مصدر إثبات الحركة")
    sources = [r[i] for r in ws.iter_rows(min_row=2, values_only=True)]
    assert sources and str(sources[0]).startswith("السلسلة"), sources[0]


def test_an_anchor_declares_the_paper_and_a_chain_row_declares_the_chain():
    from tools.to_xlsx import _assertion_source

    anchor = _assertion_source(
        "حركة — مرساة بعد ورقة غائبة (المبلغ مطبوع ولا تُقفله السلسلة)", {})
    assert anchor.startswith("الورق المطبوع"), anchor
    chain = _assertion_source("حركة مثبتة بالسلسلة", {"reread": True})
    assert chain.startswith("السلسلة") and "إعادة قراءة محكَّمة" in chain, chain
    assert _assertion_source("سطر ملخّص/افتتاحي — ليست حركة", {}) == ""


def test_the_guide_sheet_declares_what_the_file_does_not_witness(tmp_path):
    """نصف الحارس الثاني: ما لا يشهد به الملف يُكتب في الملف، ويُحرَس بالاختبار.

    الملف يشهد على نقل الصور وحسابها، ولا يشهد أن الصور صادرة عن المصرف — ولا
    شاهد في المنظومة يمسّ أصل الورق. سطرٌ يُحذف بسهولة إن لم يحرسه اختبار.
    """
    out = tmp_path / "export.xlsx"
    build(_synthetic_run(tmp_path), out, None, profile=PROFILE)
    ws = load_workbook(out)["كيف تُقرأ هذه الأوراق"]
    notes = {str(r[0]): str(r[1]) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
    assert "ما لا يشهد به هذا الملف" in notes
    assert "لا يشهد أن الصور صادرة عن المصرف" in notes["ما لا يشهد به هذا الملف"]
    assert "مصدر إثبات كل حركة" in notes, "العمود الجديد يُشرح في ورقة القراءة"


def test_markdown_renderer_keeps_rtl_and_tables():
    html_out = render("# عنوان\n\n| أ | ب |\n|---|---|\n| ١ | ٢ |\n", "عنوان")
    assert 'dir="rtl"' in html_out and 'lang="ar"' in html_out
    assert html_out.count("<th>") == 2 and html_out.count("<td>") == 2
    assert "<h1>عنوان</h1>" in html_out


def test_markdown_renderer_escapes_html_but_keeps_bold():
    html_out = render("**عريض** و<script>alert(1)</script>\n", "x")
    assert "<strong>عريض</strong>" in html_out
    assert "<script>" not in html_out
