"""ضابطُ قبول زرّ Excel — يكتبه المدقّق ويشغّله (قرارُ المالك 2026-10-07 · مراجعة ٩٦ · R96-1).

**العطبُ الذي يقيسه:** ضوابطُ الزرّ السابقة كلُّها بنت صفوفَها من تشغيلةٍ اصطناعيّة (`_load(run)`)، وهي
تحمل رقمَ الصفّ وحقولَ العقد كاملة. والتطبيقُ يحفظ صفوفًا بتسعة حقول (`app.py`). فمرّت الضوابطُ
**والزرُّ لم يُنتج ملفًّا في التطبيق قطّ** (`IntegrityError … row_no` · ثمّ `KeyError: None`).

**فهذا الضابطُ لا يبني صفًّا بيده:** يشغّل **مسارَ القراءة في `app.py` نفسَه** على العيّنة المصطنعة
المتتبَّعة (`data/sample/statement_sample.pdf`). والمستبدَلُ **نداءاتُ النموذج وحدها** (فحصُ المصرف ·
التذييل · الصفوف · البيان) بردودٍ مكتوبة، **والشبكةُ مغلقة** ⇒ $0 مُثبتًا لا مُدَّعى. ثمّ ينادي
**دالّةَ الزرّ نفسَها** (`export_xlsx`) على الحالة الناتجة، ويفتح الملفّ ويقرأ أوراقه.

**وما لا يقيسه:** النقرةَ في المتصفّح، وقراءةَ نموذجٍ حقيقيّ — تلك في البروفة المدفوعة (`REQ-93-02`).
"""

from __future__ import annotations

import socket
from decimal import Decimal
from pathlib import Path

import pytest

pytest.importorskip("gradio")
pytest.importorskip("langchain")
pytest.importorskip("openpyxl")

ROOT = Path(__file__).resolve().parents[1]
SAMPLE = ROOT / "data" / "sample" / "statement_sample.pdf"

# **ردودُ النموذج المكتوبة** — صفحتان، والسلسلةُ تُغلق، والتذييلُ تراكميّ كما في الورق (قيمٌ اصطناعيّة).
D = Decimal
_ROWS = {
    1: [{"date": "01/01/2023", "desc": "رصيد افتتاحي", "movement": None, "balance": D("500.00")},
        {"date": "05/01/2023", "desc": "شراء", "movement": D("20.00"), "balance": D("480.00")},
        {"date": "09/01/2023", "desc": "إيداع", "movement": D("35.00"), "balance": D("515.00")}],
    2: [{"date": "12/01/2023", "desc": "رصيد سابق", "movement": None, "balance": D("515.00")},
        {"date": "15/01/2023", "desc": "رسوم", "movement": D("15.00"), "balance": D("500.00")}],
}
_FOOTERS = {1: (D("20.00"), D("35.00"), D("515.00")), 2: (D("35.00"), D("35.00"), D("500.00"))}


def _page_of(image_path) -> int:
    """`…/pg-01.png` ⇒ 1 — رقمُ الصفحة من اسم الصورة التي يمرّرها التطبيق."""
    return int(Path(str(image_path)).stem.split("-")[-1])


@pytest.fixture
def read_sample(monkeypatch):
    """يشغّل مسارَ القراءة في `app.py` على العيّنة المصطنعة — والنموذجُ وحده مستبدَل، والشبكةُ مغلقة."""
    import app
    from statement_qa.footer_oracle import FooterReading

    def _no_network(*_a, **_k):
        raise AssertionError("نداءُ شبكةٍ في ضابطٍ وعدُه $0 — نداءُ نموذجٍ لم يُستبدَل")

    monkeypatch.setattr(socket.socket, "connect", _no_network)

    def probe_bank(_img, stats=None):
        return {"verdict": "alrajhi", "bank": "مصرف الراجحي"}

    def read_footer(img, stats=None):
        dr, cr, bal = _FOOTERS[_page_of(img)]
        return FooterReading(debits=dr, credits=cr, balance=bal)

    def read_rows_vlm(img, stats=None):
        pg = _page_of(img)
        if stats is not None:
            stats["page_no"] = pg
        return [dict(r) for r in _ROWS[pg]]

    def read_page_descriptions(*_a, **_k):
        raise ValueError("البيانُ مستبدَل في الضابط — يبقى بيانُ قراءة الصفحة")

    monkeypatch.setattr(app, "probe_bank", probe_bank)
    monkeypatch.setattr(app, "read_footer", read_footer)
    monkeypatch.setattr(app, "read_rows_vlm", read_rows_vlm)
    monkeypatch.setattr(app, "read_page_descriptions", read_page_descriptions)
    # **وفهرسُ البحث خارج سؤال الزرّ:** يُحمّل نموذجَ تضمينٍ قد يسأل الشبكة (أمسكه قفلُ الشبكة هنا)،
    # والزرُّ لا يقرؤه ⇒ يُستبدَل فلا يبقى في الضابط إلّا ما يبني الصفوفَ والملفّ.
    monkeypatch.setattr(app, "build_index", lambda _chunks: None)
    app.STATE.clear()
    app._process_pdf_locked(str(SAMPLE), progress=lambda *_a, **_k: None)
    yield app
    app.STATE.clear()


def test_the_read_path_holds_the_rows_the_app_builds(read_sample) -> None:
    """**المُدخَلُ صفوفُ التطبيق لا صفوفُ تشغيلة:** القراءةُ أنتجت حالةً، وصفوفُها من كود `app.py`."""
    rows = read_sample.STATE.get("rows") or []
    assert len(rows) == sum(len(v) for v in _ROWS.values()), f"عددُ الصفوف: {len(rows)}"
    assert {r.get("page") for r in rows} == {1, 2}


def test_the_download_button_builds_a_workbook_from_the_app_state(read_sample) -> None:
    """**الدعوى الحاكمة (R96-1):** زرُّ التنزيل على حالة التطبيق ⇒ ملفٌّ بأوراقه السبع، بلا «None»."""
    from openpyxl import load_workbook

    path, _note = read_sample.export_xlsx()
    assert path, "الزرُّ لم يُنتج ملفًّا على حالة التطبيق"
    wb = load_workbook(path)
    assert len(wb.sheetnames) == 7, wb.sheetnames
    moves = wb["الحركات"]
    assert moves.max_row - 1 == len(read_sample.STATE["rows"]), "صفوفُ الحركات لا تساوي صفوفَ الحالة"
    for row in wb["كيف تُقرأ هذه الأوراق"].iter_rows(values_only=True):
        assert "None" not in " ".join(str(v) for v in row if v is not None), row


def test_the_ledger_buttons_build_refresh_and_answer(read_sample, tmp_path, monkeypatch) -> None:
    """**زرّا الدفتر (R98-1 · مراجعة ٩٨):** «ابنِ الدفتر» ثمّ «حدّث» ثمّ سؤالٌ ثابت — على حالة التطبيق نفسِها.

    سقطت الثلاثةُ في كلّ ضغطةٍ منذ 10-04 (`_ledger_ui.pages_of` لا وجودَ له)، ولم يضغطها اختبارٌ عبر
    التطبيق. والدفترُ يُكتب في مجلّدٍ مؤقّت لا في `data/` (`_APP_ROOT` مُستبدَل).
    """
    app = read_sample
    monkeypatch.setattr(app, "_APP_ROOT", tmp_path)
    msg, status = app._ledger_build_click()
    assert msg and status, "زرُّ «ابنِ الدفتر» لم يُعِد رسالةً ولا حالة"
    assert list(tmp_path.rglob("*.sqlite")), f"لم يُكتب دفتر: {msg}"
    assert app._ledger_refresh_click(), "زرُّ «حدّث» لم يُعِد حالة"
    answer = app._ledger_ask_click("تغطيةُ الكشف (صفوف · صفحات · بلا إثبات)", "")
    assert answer and "تعذّر" not in answer, answer[:200]
