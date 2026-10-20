"""بوّابةُ القاعدة في التطبيق (R130): القراءةُ تكتبها، والحذفُ يمحوها، والصورُ لا تبقى في النظام.

**العِلّة المقيسة (مراجعة ١٣٠):** «القراءةُ تكتب قاعدةً» كانت **زرًّا منفصلًا** لا يُضغط؛ و«الملفّاتُ
المؤقّتةُ لا تُحذف» كانت `mkdtemp` بلا `finally` (١٩٢٥ مجلّدًا · ٧٥٥MB يوم الجلسة). والضوابطُ هنا
تشغّل **مسارَ القراءة نفسَه** على العيّنة المصطنعة (نموذجٌ مستبدَل، شبكةٌ مغلقة ⇒ $0) ثمّ تقيس:

- القاعدةَ موجودةٌ بعد المعالجة **بلا زرّ**، وفيها الصفوفُ والتذييلات.
- أسطحَ السؤال تُعاد من القاعدة بعد فقدِ الحالة (كما بعد إعادة التشغيل).
- زرَّ الحذف يترك **صفرَ ملفّ** ويُفرّغ الحالة.
- تنزيلَ المحادثة يقرأ من `qa`/`qa_evidence` لا من الرام.
- وصفرَ مجلّدٍ `rajhi_*` بعد قراءةٍ وتنزيلٍ، و`delete_cache` يحذف ملفّاتِ الحافظة بحسب المكتبة.
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

import pytest

pytest.importorskip("gradio")
pytest.importorskip("langchain")
pytest.importorskip("langgraph.checkpoint.sqlite")

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from statement_qa import ledger as L  # noqa: E402

#: العيّنةُ والـfixture من ملفّ ضابط الزرّ (نموذجٌ مستبدَلٌ وشبكةٌ مغلقة) — تُطلَب **بالاسم** عبر
#: `request` فلا يُعاد تعريفُ اسمٍ فيُسقطه الحارسُ البنيويّ (درسُ R57-4)، ولا تُنسخ استبدالاتُها.
from test_button_builds_a_workbook import SAMPLE, read_sample  # noqa: E402


@pytest.fixture
def app(request):
    """حالةُ التطبيق بعد قراءةِ العيّنة — الـfixture الأصلُّ `read_sample` بلا إعادة تعريفٍ هنا."""
    return request.getfixturevalue(read_sample.__name__)


def _db(app) -> Path:
    return app._ledger_ui.ledger_file(app.DATA_HOME, app._ledger_key())


def test_processing_writes_the_database_without_pressing_a_button(app):
    path = _db(app)
    assert path.exists(), "القراءةُ لم تكتب القاعدةَ (قرارُ المالك: عند انتهاء المعالجة لا بزرّ)"
    conn = L.read_ledger(path)
    try:
        rows = L.read_rows(conn)
        footers = L.read_footers(conn)
    finally:
        conn.close()
    assert len(rows) == len(app.STATE["rows"]) == 5, len(rows)
    assert set(footers) == {1, 2}
    assert footers[1]["debits"] and footers[1]["verdict"] == "ok", footers[1]


def test_the_question_surfaces_are_rebuilt_from_the_database_alone(app, monkeypatch):
    """**قبولُ القرار:** «الأداةُ تجيب من القاعدة ولو فرغت الذاكرة» — تُمسح الحالةُ ثمّ تُبنى منها."""
    key = app._ledger_key()
    app.STATE.clear()
    # فهرسُ البحث خارج الدعوى (الـfixture يستبدله بـNone لئلّا يحمّل نموذجَ تضمين) — يُثبت هنا.
    monkeypatch.setattr(app, "build_index", lambda chunks: f"index:{len(chunks)}")
    assert app._load_from_ledger(key), "تعذّرت إعادةُ البناء من القاعدة"
    assert app.STATE.get("store"), "لم يُبنَ الفهرسُ من القاعدة"
    assert len(app.STATE["rows"]) == 5
    assert app.STATE.get("footers"), "التذييلاتُ لم تُقرأ من القاعدة"


def test_deleting_a_statement_leaves_zero_files_and_empties_the_state(app):
    path = _db(app)
    folder = path.parent
    assert path.exists()
    out = app._delete_statement_click()
    assert "حُذفت" in out[0], out[0]
    assert not list(folder.glob("*.sqlite*")), list(folder.iterdir())
    assert not (app.DATA_HOME / "active.json").exists(), "مؤشّرُ الكشف بقي بعد الحذف"
    assert app.STATE == {}, "الحالةُ لم تُفرَّغ بعد الحذف"


def test_the_conversation_downloads_from_the_database(app):
    """**البند ١٠:** الملفُّ يُبنى من `qa`+`qa_evidence` (لا من ذاكرةٍ حيّة) — فيبقى بعد الإغلاق."""
    path = _db(app)
    conn = L.open_ledger(path)                # الكتابةُ تحتاج فتحًا كاملًا (read_ledger للقراءة وحدها)
    try:
        ids = L.row_ids_for(conn, [(1, 2), (1, 3)])
        L.record_qa(conn, "2026-10-20T09:00:00", "كم مجموع السحوبات؟",
                    "المجموع 20.00 (صفحة 1، صف 2)", ids)
    finally:
        conn.close()
    app.STATE.clear()                         # كأنّ التطبيق أُعيد تشغيلُه
    got, note = app._conversation_export()
    assert got and Path(got).exists(), note
    from openpyxl import load_workbook

    wb = load_workbook(got)
    chat = list(wb["المحادثة"].values)
    ev = list(wb["الأدلّة"].values)
    assert chat[1][2] == "كم مجموع السحوبات؟" and "المجموع 20.00" in chat[1][3], chat
    assert ev[1][1:3] == (1, 2), ev                 # الصفّ الأولُ من الأدلّة: صفحة 1، صف 2


def test_no_rajhi_work_folder_is_left_in_the_system_temp_dir(app):
    """**قبولُ R130-2:** بعد قراءةٍ وتنزيلٍ ⇒ صفرُ مجلّدٍ `rajhi_*` في مجلّد النظام."""
    tmp = Path(tempfile.gettempdir())
    before = {p.name for p in tmp.glob("rajhi_*")}
    app._process_pdf_locked(str(SAMPLE), progress=lambda *_a, **_k: None)
    out, _note = app.export_xlsx()
    assert out and Path(out).exists()
    after = {p.name for p in tmp.glob("rajhi_*")}
    assert after - before == set(), f"مجلّداتٌ بقيت: {sorted(after - before)}"


def test_the_library_purges_the_cache_files_it_created(app):
    """**«بمكوّن المكتبة لا منظّفٍ يدويّ»:** `delete_files_created_by_app` (التي يناديها `delete_cache`
    وعند الإغلاق) تحذف ملفّاتِ المخرجات التي سجّلتها المكوّنات."""
    from gradio.route_utils import delete_files_created_by_app

    app.export_xlsx()
    tracked = [f for s in app.demo.temp_file_sets for f in s if Path(f).exists()]
    assert tracked, "لم تُسجَّل ملفّاتُ الحافظة (النسخُ لم يقع؟)"
    delete_files_created_by_app(app.demo, age=None)
    left = [f for f in tracked if Path(f).exists()]
    assert not left, f"ملفّاتٌ لم تُحذف: {left[:3]}"


def test_the_table_filter_uses_a_chosen_column_and_a_date_range(app):
    """**البندان ٧ و٨:** الفلترةُ بعمودٍ مختار («النوع») وبمدى تاريخٍ من/إلى.

    والمدى يُقاس على صفوفٍ بتواريخ ثمانيةِ أرقام (كما تُطبع في كشوف الراجحي) — لا على تواريخ
    العيّنة المصطنعة (`01/01/2023` بصيغة يوم/شهر/سنة، وهي شكلٌ **لا يقرؤه** قارئُ العقد المعلَن؛
    فلا يُخترع لها قارئٌ ثالث في ضابط).
    """
    everything = app.filter_rows()
    typed = app.filter_rows("النوع", "إيداع")
    assert not typed.empty and len(typed) < len(everything)
    assert all("إيداع" in str(v) for v in typed["النوع"]), typed["النوع"].tolist()
    app.STATE["rows"] = [
        {"page": 1, "kind": "txn", "balance": "10.00", "movement": "10.00",
         "side": "debit", "ok": True, "desc": "أ", "date": "20260301", "type": "سحب"},
        {"page": 1, "kind": "txn", "balance": "20.00", "movement": "10.00",
         "side": "debit", "ok": True, "desc": "ب", "date": "20260415", "type": "سحب"},
        {"page": 1, "kind": "txn", "balance": "30.00", "movement": "10.00",
         "side": "debit", "ok": True, "desc": "ج", "date": "", "type": "سحب"},
    ]
    dated = app.filter_rows("", "", "2026-03-01", "2026-03-31")
    assert len(dated) == 1 and dated.iloc[0]["الوصف"] == "أ", dated
    assert len(app.filter_rows("", "", "", "2026-03-31")) == 1, "حدُّ «إلى» لم يُطبَّق"
    assert len(app.filter_rows("", "", "2026-04-01")) == 1, "حدُّ «من» لم يُطبَّق"


def test_rebuilding_the_ledger_after_a_restart_does_not_wipe_the_rows(app):
    """**العطب P0 الذي اتّفق عليه المقاعد الثلاثة:** `_load_from_ledger` كان يضع صفوفَ القاعدة
    (`pg` بلا `page`) في `rows_contract`، وزرُّ «أعد بناء الدفتر» يُعيد إدخالها ⇒ `dropped_no_page`
    والصفوفُ كلُّها تُقلَّم (قِيس: ٦ ⇒ ٠) والرسالةُ تقول «بُني ✅». فالسطحُ يقرأ الآن `page` من
    `read_rows`، وهذا الضابطُ يفشل على الشيفرة قبل الإصلاح.
    """
    key = app._ledger_key()
    path = _db(app)
    conn = L.read_ledger(path)
    try:
        before = len(L.read_rows(conn))
    finally:
        conn.close()
    assert before == 5

    app.STATE.clear()                          # إعادةُ تشغيل
    assert app._load_from_ledger(key)
    msg, _status = app._ledger_build_click()   # الزرُّ الذي كان يمحو

    conn = L.read_ledger(path)
    try:
        after = len(L.read_rows(conn))
    finally:
        conn.close()
    assert after == before, f"أعاد البناءُ كتابةً محت؟ {before} ⇒ {after} — الرسالة: {msg}"
    assert "0 صفًّا" not in msg, msg


def test_pages_are_counted_from_the_database_not_from_rows(app):
    """**عددُ الصفحات بعد إعادة التشغيل (مقعدُ المواصفة P2):** ٦٢٦ صفًّا لا تعني ٦٢٦ صفحة — الصفحةُ
    بلا صفٍّ (غلاف/ملخّص) تبقى في جدول `pages`، و`n_pages` من الصفوف كان يقول «ينتهي عند 626»."""
    key = app._ledger_key()
    app.STATE.clear()
    assert app._load_from_ledger(key)
    assert app.STATE["n_pages"] == 2, app.STATE["n_pages"]


def test_the_deleted_statement_leaves_no_file_at_all(app):
    """**«صفرُ ملفّ» بالمعنى:** يُعدّ **كلُّ** ما في مجلّد البيانات — لا `*.sqlite*` وحدها (كان
    `job.lock` يبقى فيُكذّب الجملة)."""
    _db(app)
    app._delete_statement_click()
    left = sorted(p.name for p in app.DATA_HOME.rglob("*") if p.is_file())
    # **«صفرُ ملفّ» يُعرَّف على بيانات الكشف (R131-5):** `job.lock` ملفُّ تنسيقٍ بين النوافذ
    # (رقمُ العملية ووقتها) لا بياناتُ كشف — ويبقى لأنّ حذفَه أثناء تشغيلٍ يفتح البابَ لمهمّتين.
    assert left == ["job.lock"], f"بقيت ملفّاتُ كشفٍ بعد الحذف: {left}"
    assert not list(app.DATA_HOME.rglob("*.sqlite*")) and not (app.DATA_HOME / "active.json").exists()


def test_the_screen_number_and_the_tool_number_are_the_same(app):
    """**رقمُ الصفّ الواحد (مقعدُ البنية P1):** سطحُ العرض كان يُرقّم 1..n والأدواتُ 1..m، فكلُّ صفٍّ
    بلا رصيدٍ يُزيحهما. والآن `seq` (موضعُ الصفّ في العقد) هو رقمُ الاثنين."""
    from statement_qa.qa_tools import make_qa_tools

    rows = app.STATE["rows"]
    assert [r["seq"] for r in rows] == sorted(r["seq"] for r in rows)
    struct = [{"pg": 1, "row_no": 1, "page": 1, "balance": "10.00", "movement": None,
               "side": "", "opening": True, "desc": "افتتاحي", "date": "2026-03-01"},
              {"pg": 1, "row_no": 2, "page": 1, "balance": None, "movement": None,
               "side": "", "opening": False, "desc": "سطرُ ملخّص", "date": ""},
              {"pg": 1, "row_no": 3, "page": 1, "balance": "8.00", "movement": "2.00",
               "side": "debit", "opening": False, "desc": "سحب", "date": "2026-03-02"}]
    conn = L.open_ledger(app.DATA_HOME / "numb.sqlite")
    try:
        L.ingest(conn, struct, {1: {"verdict": "ok"}})
        got = {t.name: t for t in make_qa_tools(conn=conn)}
        text = got["search_rows"].invoke({"keyword": "سحب"})
        assert "صف 3" in text, text          # موضعُه في العقد (الثالث) لا الثاني في العرض
    finally:
        conn.close()
    view = app._view_rows(struct)
    assert [r["seq"] for r in view] == [1, 3], view


def test_an_empty_filter_keeps_the_columns(app):
    """الفلترةُ إلى صفر تُبقي رؤوسَ الجدول (وإلا قُرئت عطبًا لا نتيجةً فارغة)."""
    empty = app.filter_rows("النوع", "لا_يوجد_هذا_أبدًا")
    assert empty.empty
    assert list(empty.columns)[:3] == ["#", "الصفحة", "التاريخ"], list(empty.columns)


def test_the_filter_columns_match_the_rendered_record():
    """قائمةُ الفلترة تُقابَل بمفاتيح السجلّ المعروض — إعادةُ تسميةِ عمودٍ تُسقط هذا الضابط."""
    import app as app_module
    from statement_qa.render import row_record

    keys = list(row_record(1, {"page": 1, "kind": "txn", "balance": "1", "movement": "1",
                               "side": "debit", "ok": True, "desc": "x", "date": ""}).keys())
    assert list(app_module.FILTER_COLUMNS[1:]) == keys, (app_module.FILTER_COLUMNS, keys)


def test_the_app_writes_the_question_and_the_intent_reads_it(app):
    """**قبولُ القرار: «قصدُ الأسئلة المسجَّلة يعرض ما سُئل فعلًا».** الضابطُ السابق كان يُدرج صفَّ
    `qa` بيده ثمّ يقيس القراءة — فلو حُذف كاتبُ التطبيق بقي أخضر. هذا يُنادي الكاتبَ نفسَه
    (`_record_qa_row`) ثمّ يقرأ بالقصد، ويقيس الترجمةَ من رقم الصفّ إلى `(صفحة، صف)` في القاعدة.
    """
    from types import SimpleNamespace

    res = SimpleNamespace(used_row_nos=[2], cited_row_ids=[], refused=False)
    app._record_qa_row(res, "كم مجموع السحوبات؟", "المجموع 20.00 (صفحة 1، صف 2)")
    got = app._ledger_ui.ask(app.DATA_HOME, app._ledger_key(), "الأسئلةُ المسجَّلة")
    flat = " | ".join(str(c) for row in got["rows"] for c in row)
    assert "كم مجموع السحوبات؟" in flat, flat
    assert "10.00" in flat or "صفحة 1" in flat, flat      # دليلُ الصفّ وصل من القاعدة لا من الرام


# ————— R131-10 (مراجعة ١٣١): الحذفُ يعدّ نسخَ العرض في حافظة Gradio ولا يبتلع فشلَها —————
# **العِلّة المقيسة:** كان `delete_files_created_by_app` في `try/except Exception: pass`، ورسالةُ الحذف
# تعدّ ملفّاتِ القاعدة وحدَها ⇒ «صفرُ ملفّ» عن نسخ العرض (صورُ الصفحات · الإكسل · المحادثة المنزَّلة،
# وكلُّها بياناتٌ حقيقيّة) دعوى بلا عدّ، وفشلُ حذفها صامت. والضابطان يقيسان العدَّ والفشلَ معًا.

def _display_copy(app, tmp_path) -> Path:
    """نسخةُ عرضٍ في حافظة Gradio **بمكوّن المكتبة نفسِه** الذي يستعمله التطبيق (`_cache_path`)."""
    src = tmp_path / "conversation.xlsx"
    src.write_text("نسخةُ عرض", encoding="utf-8")
    return Path(app._cache_path(app.conv_file, src))


def test_the_delete_counts_the_display_copies_it_removed(app, tmp_path):
    kept = _display_copy(app, tmp_path)
    # **والعددُ من قائمة المكتبة قبل الضغط:** قراءةُ العيّنة وضعت صورَ صفحاتها في الحافظة (معرضُ الصفحات)،
    # وهي ممّا يُعدّ — فالمتوقَّعُ «كلُّ ما في القائمة» لا رقمٌ مكتوبٌ باليد.
    n = len(app._display_copies())
    assert kept.exists() and n >= 2, n               # صورُ القراءة + المحادثةُ المنزَّلة
    msg = app._delete_statement_click()[0]
    assert not kept.exists(), "نسخةُ العرض بقيت بعد الحذف"
    assert f"نسخُ العرض المؤقّتة: حُذف **{n}** · بقي **0**" in msg, msg
    assert app._display_copies() == set()


def test_a_failed_display_cleanup_is_named_not_swallowed(app, tmp_path, monkeypatch):
    import gradio.route_utils as ru

    kept = _display_copy(app, tmp_path)
    n = len(app._display_copies())

    def _refuse(*_a, **_k):
        raise PermissionError("محاكاةُ رفضِ النظام")

    monkeypatch.setattr(ru, "delete_files_created_by_app", _refuse)
    msg = app._delete_statement_click()[0]
    assert kept.exists()
    assert "تعذّر حذفُ نسخ العرض المؤقّتة (PermissionError)" in msg, msg
    assert f"حُذف **0** · بقي **{n}**" in msg, msg


def test_the_conversation_export_keeps_newlines_and_pipes_exactly(app):
    """**قبولُ R131-6:** وصفٌ فيه سطرٌ جديدٌ و`|` يعود من الملفّ **حرفًا بحرف** — الجدولُ اليدويّ
    كان يبدّل `|` بـ`/` ويسقط السطرَ الجديد، فينكسر الصفُّ في ٣٨٨٣ وصفًا من ٥٨٠٩ في كشف المالك.
    """
    from openpyxl import load_workbook

    path = _db(app)
    conn = L.open_ledger(path)
    try:
        ids = L.row_ids_for(conn, [(1, 2)])
        row_id = ids[0] if ids else conn.execute("SELECT id FROM rows_verified LIMIT 1").fetchone()[0]
        conn.execute("UPDATE rows_verified SET desc = ? WHERE id = ?",
                     ("سطرٌ أول\nسطرٌ ثانٍ | عمود", row_id))
        conn.commit()
        L.record_qa(conn, "2026-10-20T09:30:00", "سؤالُ السطرين", "جواب", [row_id])
    finally:
        conn.close()
    got, note = app._conversation_export()
    assert got, note
    wb = load_workbook(got)
    descs = [r[3] for r in wb["الأدلّة"].values]
    assert "سطرٌ أول\nسطرٌ ثانٍ | عمود" in descs, descs


def test_the_delete_is_refused_while_a_read_holds_the_job_lock(app):
    """**R131-5 (P2):** الحذفُ يجري تحت قفل المهمّة نفسِه. كان يزيل ملفَّ القفل بلا شرط ⇒ مَن بعده
    يأخذه وقراءةٌ جاريةٌ تعمل ⇒ مهمّتان مدفوعتان، والقراءةُ الجاريةُ تُعيد كتابةَ البيانات بعد «حُذفت».
    """
    from statement_qa.job_lock import job_lock

    path = _db(app)
    assert path.exists()
    with job_lock(app.LOCK_PATH):                 # قراءةٌ جارية
        out = app._delete_statement_click()
        assert "لم يُحذف شيء" in out[0], out[0]
        assert path.exists(), "حُذفت القاعدةُ رغم تشغيلٍ جارٍ"
    out = app._delete_statement_click()           # بعد تحرير القفل
    assert "حُذفت" in out[0], out[0]
    assert not path.exists()


def test_the_three_surfaces_agree_after_an_incomplete_reread(app):
    """**قبولُ R131-3 (الضابطُ الثلاثيّ الذي تبنّاه المدقّق):** بعد إعادة قراءةٍ تعثّرت فيها صفحة،
    **الشاشة = القاعدة = الأدوات** و`seq` متّصل. يحاكي: تُبقى صفحةٌ واحدةٌ في القاعدة (كما يفعل
    `ingest` بعد تعثّر الأخرى)، ثمّ تُعاد الأسطحُ من القاعدة ويُقاس الثلاثةُ معًا.
    """
    from statement_qa.qa_tools import make_qa_tools

    key = app._ledger_key()
    path = _db(app)
    conn = L.open_ledger(path)
    try:
        rows = L.read_rows(conn)
        pages = L.read_footers(conn)
        keep = [r for r in rows if r["page"] == 1]
        L.ingest(conn, keep, {1: pages[1]})        # الصفحةُ ٢ «تعثّرت»: لا صفوفَ لها الآن
    finally:
        conn.close()

    app.STATE.clear()                              # إعادةُ تشغيل
    assert app._load_from_ledger(key)
    screen = [int(r["seq"]) for r in app.STATE["rows"]]
    number_of_displayed = len(app._rows_df(app.STATE["rows"]))

    conn = L.read_ledger(path)
    try:
        db = L.read_rows(conn)
        tools = {t.name: t for t in make_qa_tools(conn=conn)}
        counted = tools["count_movements"].invoke({})
    finally:
        conn.close()
    db_seq = [int(r["seq"]) for r in db]
    assert screen == db_seq, f"الشاشة {screen} ≠ القاعدة {db_seq}"
    assert db_seq == list(range(1, len(db_seq) + 1)), db_seq
    assert number_of_displayed == len(db_seq)
    # الأداةُ تعدّ **الحركات** (لا صفَّ الرصيد الافتتاحيّ) — فالمقابلةُ بدلالتها لا بعدد الصفوف.
    movements = sum(1 for r in db if not r.get("opening") and r.get("movement") is not None)
    assert f"العدد = {movements}" in counted, counted
