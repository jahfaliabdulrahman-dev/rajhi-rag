#!/usr/bin/env python3
"""اختبار الحقن: **كل فحص بوابة يُثبَت بأنه يفشل على ملف مسموم.**

الدرس الذي جاء من مراجعة خارجية ثلاث مرات: الفحص الميت صنفان، ولكلٍّ كاشفٌ مختلف:

  ١) **ميتٌ بمفتاح خاطئ** — يقرأ حقلاً لا وجود له فيُصفّر مرشِّحه. يكشفه **إعلان العدد**
     (الفحص الحيّ يقول «من N»)، ولا يكشفه أنه «مرّ».
  ٢) **ميتٌ بشرطٍ مُعطَّل** — يقول العدد، ويُظهر النقص، **ويمرّ** (`or not required`).
     ولا يكشفه إلا **الحقن**: يُفسَد شيءٌ يجب أن يُفشل، فإن مرّ فحصٌ على سمّه فهو ميت.

    python3 tools/injection_test.py --xlsx ~/Downloads/rajhi-rag-export-629p.xlsx

يُرجع 0 إن عضّ الجميع، وغير صفر إن مرّ فحصٌ على سمّه.
"""
from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

PROJ = Path(__file__).resolve().parent.parent


def run_gate(run: str, xlsx: Path) -> tuple[int, str]:
    p = subprocess.run([sys.executable, str(PROJ / "tools/verify_close.py"),
                        "--run", run, "--xlsx", str(xlsx)],
                       cwd=PROJ, capture_output=True, text=True)
    return p.returncode, (p.stdout + p.stderr)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--run", default="data/local_sample/slice_629p")
    args = ap.parse_args()
    src = Path(args.xlsx).expanduser()
    if not src.exists():
        print(f"الملف غير موجود: {src}")
        return 2
    import openpyxl

    results: list[tuple[str, bool, str]] = []

    code, out = run_gate(args.run, src)
    clean_ok = code == 0
    results.append(("الملف السليم يجب أن يمرّ", clean_ok, f"exit={code}"))

    poisons = {
        "حذف تاريخ حركة مثبتة": ("التاريخ (ميلادي)", None),
        "تناقض «كما طُبعت» مع المثبتة": ("الحركة كما طُبعت", "٩٩٩٩٩.٩٩"),
    }
    for label, (col, value) in poisons.items():
        tmp = Path(tempfile.mkdtemp(prefix="inject-")) / "poisoned.xlsx"
        shutil.copy(src, tmp)
        wb = openpyxl.load_workbook(tmp)
        ws = wb["الحركات"]
        h = [c.value for c in ws[1]]
        iv, it = h.index("حكم السلسلة على الصفّ") + 1, h.index(col) + 1
        hit = 0
        for r in range(2, ws.max_row + 1):
            if str(ws.cell(r, iv).value or "").startswith("حركة") and ws.cell(r, it).value not in (None, ""):
                ws.cell(r, it).value = value
                hit += 1
                break                      # **صفٌّ واحد** — أدقّ من تعميم السمّ
        wb.save(tmp)
        code, out = run_gate(args.run, tmp)
        died = code != 0 and "FAIL" in out
        results.append((label, hit > 0 and died, f"خلايا مسمومة={hit} exit={code}"))

    print("── اختبار الحقن ──")
    for label, ok, detail in results:
        print(f"  {'✓' if ok else '✗'} {label} — {detail}")
    bad = [l for l, ok, _ in results if not ok]
    if bad:
        print(f"\nBLOCK — فحصٌ مرّ على سمّه: {bad}")
        return 1
    print("\nPASS — كل فحص عضّ سمّه، والملف السليم مرّ.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
