import subprocess, sys, shutil  # noqa (نصّ المدقّق حرفيّ — الوسم لإسكات استيرادين غير مستعملين فيه)
from openpyxl import load_workbook
SRC='~/Downloads/rajhi-rag-export-629p.xlsx'
OUT=sys.argv[1]
def load():
    wb=load_workbook(SRC); ws=wb['الحركات']; h=[c.value for c in ws[1]]
    return wb,ws,h
def col(h,n): return h.index(n)+1
def first_proven(ws,h):
    for r in range(2,ws.max_row+1):
        if str(ws.cell(r,col(h,'حكم السلسلة على الصفّ')).value or '').startswith('حركة مثبتة'):
            return r
def p_negative(wb,ws,h):
    r=first_proven(ws,h); c=col(h,'مدين')
    if ws.cell(r,c).value is None: c=col(h,'دائن')
    ws.cell(r,c).value=-abs(float(ws.cell(r,c).value))
def p_bothcols(wb,ws,h):
    r=first_proven(ws,h); ws.cell(r,col(h,'مدين')).value=1; ws.cell(r,col(h,'دائن')).value=1
def p_tanbih(wb,ws,h):
    ws.cell(first_proven(ws,h),col(h,'تنبيه')).value='تناقض وصف↔اتجاه'
def p_gapverdict(wb,ws,h):
    for r in range(2,ws.max_row+1):
        if 'قيد فجوة' in str(ws.cell(r,col(h,'حكم السلسلة على الصفّ')).value or ''):
            ws.cell(r,col(h,'حكم السلسلة على الصفّ')).value='قيد فجوة مسح — غير مُثبت'; return
def p_dropdate(wb,ws,h):
    ws.cell(first_proven(ws,h),col(h,'التاريخ (ميلادي)')).value=None
def p_sheetname(wb,ws,h): wb['ما لم يُثبت'].title='ما لم يُثبتX'
def p_identity(wb,ws,h):
    m=wb['الملخص']
    for row in m.iter_rows():
        if row[0].value and 'حكم الهوية' in str(row[0].value): row[1].value='مخالف'; return
def p_sumdebit(wb,ws,h):
    ws.cell(first_proven(ws,h),col(h,'مدين')).value=(ws.cell(first_proven(ws,h),col(h,'مدين')).value or 0)+1000
CASES={'قيمة سالبة':p_negative,'الرقمان معاً':p_bothcols,'تنبيه وصف↔اتجاه':p_tanbih,
       'قيد فجوة غير مُثبت':p_gapverdict,'حذف تاريخ حركة':p_dropdate,'اسم ورقة':p_sheetname,
       'حكم الهوية':p_identity,'زيادة 1000 على مدين':p_sumdebit}
name=sys.argv[2]
wb,ws,h=load(); CASES[name](wb,ws,h); wb.save(OUT)
